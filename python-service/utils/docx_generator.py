"""Render Markdown reports to Word locally; never fetch model-generated resources."""
import os
import re
import tempfile
from html.parser import HTMLParser

import markdown
from docx import Document
from docx.shared import Inches, Pt


def extract_markdown_content(text):
    match = re.match(r"^(?:'''|```)markdown\s*(.*?)(?:'''|```)\s*$", text.strip(), re.S | re.I)
    return match.group(1).strip() if match else text.strip()


def wrap_as_markdown_codeblock(text):
    return f"```markdown\n{text.strip()}\n```"


class _WordRenderer(HTMLParser):
    def __init__(self, document):
        super().__init__(convert_charrefs=True)
        self.document = document
        self.paragraph = None
        self.bold = self.italic = self.skip = 0
        self.lists = []
        self.table = None
        self.row = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.skip += 1
        if self.skip:
            return
        if tag in ('ul', 'ol'):
            self.lists.append(tag)
        elif re.fullmatch('h[1-6]', tag):
            self.paragraph = self.document.add_heading(level=int(tag[1]))
        elif tag == 'li':
            self.paragraph = self.document.add_paragraph(style='List Number' if self.lists and self.lists[-1] == 'ol' else 'List Bullet')
        elif tag in ('p', 'pre', 'blockquote'):
            self.paragraph = self.document.add_paragraph()
        elif tag in ('strong', 'b'):
            self.bold += 1
        elif tag in ('em', 'i'):
            self.italic += 1
        elif tag == 'br' and self.paragraph:
            self.paragraph.add_run().add_break()
        elif tag == 'table':
            self.table = self.document.add_table(rows=0, cols=0)
            self.table.style = 'Table Grid'
        elif tag == 'tr':
            self.row = []
        elif tag in ('td', 'th') and self.table is not None:
            self.row.append('')
        elif tag == 'img':
            # Reports are text evidence; omit URLs/files supplied by model text.
            self.handle_data(dict(attrs).get('alt', ''))

    def handle_endtag(self, tag):
        if tag in ('script', 'style') and self.skip:
            self.skip -= 1
            return
        if self.skip:
            return
        if tag in ('ul', 'ol') and self.lists:
            self.lists.pop()
        if tag in ('strong', 'b'):
            self.bold = max(0, self.bold - 1)
        if tag in ('em', 'i'):
            self.italic = max(0, self.italic - 1)
        if tag == 'tr' and self.table is not None and self.row:
            while len(self.table.columns) < len(self.row):
                self.table.add_column(Inches(1.3))
            cells = self.table.add_row().cells
            for index, value in enumerate(self.row):
                cells[index].text = value.strip()
            self.row = []
        if tag == 'table':
            self.table = None
        if tag in ('p', 'li', 'pre', 'blockquote') or re.fullmatch('h[1-6]', tag):
            self.paragraph = None

    def handle_data(self, data):
        if self.skip or not data:
            return
        if self.table is not None:
            if self.row:
                self.row[-1] += data
            return
        if self.paragraph is None:
            if not data.strip():
                return
            self.paragraph = self.document.add_paragraph()
        run = self.paragraph.add_run(data)
        run.bold, run.italic = bool(self.bold), bool(self.italic)


def generate_docx_from_markdown(markdown_input, output_path, template_path=None, is_text=False):
    from pathlib import Path
    text = markdown_input if is_text else Path(markdown_input).read_text(encoding='utf-8')
    text = extract_markdown_content(text)
    if not text.strip():
        raise ValueError('报告内容为空')
    document = Document()
    document.styles['Normal'].font.name = 'Microsoft YaHei'
    document.styles['Normal'].font.size = Pt(11)
    renderer = _WordRenderer(document)
    renderer.feed(markdown.markdown(text, extensions=['tables', 'fenced_code']))
    renderer.close()
    # Keep a previous complete export readable if a concurrent retry fails.
    directory = str(Path(output_path).resolve().parent)
    fd, temporary = tempfile.mkstemp(prefix='report-export-', suffix='.docx', dir=directory)
    os.close(fd)
    try:
        document.save(temporary)
        os.replace(temporary, output_path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return output_path
